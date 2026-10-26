"""Fill the dashboard template (and the Excel screener) from the radar, the judgments and the target prices.
Deterministic: code, never the model. The templates are FILLED, never rebuilt.

    python p2p/dashboard.py [--thesis thesis.md] [--out outputs] [--excel]

Reads outputs/radar.csv (else every company), outputs/judgments.csv (else nobody is Priority), data/market.json.
Writes outputs/p2p_dashboard.html, outputs/screener.csv and, with --excel, outputs/p2p_screener.xlsx (needs openpyxl).

Statuses: Reject rows leave the screener. Priority rows are "Priority - Actionable" when the takeout price is at or
below the target price (section 4 of the thesis), else "Priority - Not Actionable". Everything else is "Radar".
Consolidators are listed for reference under "Key Sector Public Consolidators".
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from p2p.target_price import evaluate, load_market, read_settings, rounded  # noqa: E402

COMPANIES = ROOT / "data" / "companies.csv"
TEMPLATE_HTML = ROOT / "templates" / "p2p_dashboard_template.html"
TEMPLATE_XLSX = ROOT / "templates" / "p2p_screener_template.xlsx"
SECTIONS = ["Priority - Actionable", "Priority - Not Actionable", "Radar", "Key Sector Public Consolidators"]
CSV_COLUMNS = ["section", "ticker", "name", "team", "verdict", "currency", "price", "mkt_cap_m", "ev_m", "ev_sales", "ev_ebitda",
               "net_leverage", "max_entry_price", "current_takeout", "current_irr", "multiple_price", "actionable_below", "ytd", "vs_vwap"]


def read_rows(path: Path) -> list:
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def ratio(a, b):
    if a is None or b is None or not b:
        return None
    return a / b


def build_rows(thesis: Path, out_dir: Path) -> tuple:
    companies = {r["ticker"]: r for r in read_rows(COMPANIES)}
    radar = out_dir / "radar.csv"
    tickers = [r["ticker"] for r in read_rows(radar)] if radar.exists() else list(companies)
    consolidators = [t for t, r in companies.items() if r["role"] == "consolidator" and t not in tickers]
    judgments = {}
    jpath = out_dir / "judgments.csv"
    if jpath.exists():
        judgments = {r["ticker"]: r for r in read_rows(jpath)}
    market = load_market()
    settings = read_settings(thesis)
    rows, rejected, failed = [], [], []
    for ticker in tickers + consolidators:
        c, m = companies[ticker], market.get(ticker)
        if m is None:
            continue
        j = judgments.get(ticker, {})
        verdict = (j.get("final_verdict") or j.get("verdict") or "").strip()
        tp = rounded(evaluate(ticker, market, settings))
        if c["role"] == "consolidator":
            section = "Key Sector Public Consolidators"
        elif verdict == "Reject":
            rejected.append(ticker)
            continue
        elif verdict == "Priority":
            section = "Priority - Actionable" if tp["actionable"] else "Priority - Not Actionable"
        else:
            section = "Radar"
            if verdict in ("Error", "ParseError"):  # the judge failed: shown on the radar, flagged, never silent
                failed.append(ticker)
        ebitda, revenue, ev = m.get("ebitda_ttm"), m.get("revenue_ttm"), m.get("ev")
        net_debt = None if m.get("total_debt") is None or m.get("total_cash") is None else m["total_debt"] - m["total_cash"]
        nm = ebitda is not None and ebitda <= 0
        row = {
            "section": section, "ticker": ticker, "name": f"{c['name']} ({ticker})", "team": c["team"], "currency": m["currency"],
            "mkt_cap_m": None if m.get("mkt_cap") is None else round(m["mkt_cap"] / 1e6, 3),
            "ev_m": None if ev is None else round(ev / 1e6, 3),
            "ev_sales": None if ratio(ev, revenue) is None else round(ev / revenue, 3),
            "ev_ebitda": "NM" if nm else (None if ratio(ev, ebitda) is None else round(ev / ebitda, 3)),
            "net_leverage": "NM" if nm else (None if ratio(net_debt, ebitda) is None else round(net_debt / ebitda, 3)),
            "max_entry_price": tp["max_entry_price"], "current_takeout": tp["current_takeout"], "current_irr": tp["current_irr"],
            "multiple_price": tp["multiple_price"], "actionable_below": tp["actionable_below"], "exit_margin_used": tp["exit_margin_used"],
            "price": m["price"], "date": m["date"], "vwap90": m["vwap90"], "d1": m["d1"], "w1": m["w1"], "ytd": m["ytd"], "vs_vwap": m["vs_vwap"],
            "history": m["history"], "signals": "",
            "data": "judgment failed" if ticker in failed else ("" if verdict else "no judgment yet"), "est": "",
            "src": {"price": f"close {m['date']}, {m['source']}", "financials": f"TTM, {m['source']}",
                    "target_price": f"thesis section 4 ({settings.method}), p2p/target_price.py"},
            "status_note": (f"{verdict}: {j.get('note', '')}".strip(": ") if verdict else "Not judged yet (python p2p/judge.py --radar)"),
            "lbo_steps": tp["steps"], "verdict": verdict,
        }
        rows.append(row)
    rows.sort(key=lambda r: (SECTIONS.index(r["section"]), r["name"]))
    return rows, rejected, failed, settings


def sector_name(thesis: Path) -> str:
    if not thesis.exists():
        return "All teams"
    # Only the "Your value" column counts: the Example column of a blank thesis.md must not name the sector.
    rows = [l for l in thesis.read_text(encoding="utf-8").splitlines() if l.strip().startswith("|")]
    header = next((r for r in rows if "your value" in r.lower()), None)
    team = next((r for r in rows if re.match(r"^\|\s*Team\s*\|", r)), None)
    if header and team:
        col = [c.strip().lower() for c in header.strip().strip("|").split("|")].index("your value")
        cells = [c.strip() for c in team.strip().strip("|").split("|")]
        if col < len(cells) and cells[col]:
            return cells[col]
    return "All teams"


def write_html(rows: list, sector: str, as_of: str, out: Path) -> None:
    template = TEMPLATE_HTML.read_text(encoding="utf-8")
    m = re.search(r'(<script id="screener-data"[^>]*>)(.*?)(</script>)', template, re.S)
    if not m:
        raise SystemExit("template: <script id=\"screener-data\"> block not found")
    current = json.loads(m.group(2))
    data = {"sector": sector, "as_of": as_of, "source": "data/market.json (frozen dataset, Yahoo Finance, retrieved 2026-10-19)",
            "multiples_basis": current.get("multiples_basis", "LTM"), "defs": current.get("defs", {}),
            "rows": [{k: v for k, v in r.items() if k not in ("team", "verdict")} for r in rows]}
    filled = template[:m.start(2)] + json.dumps(data, ensure_ascii=False) + template[m.end(2):]
    out.write_text(filled, encoding="utf-8")


def write_csv(rows: list, out: Path) -> None:
    with open(out, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_excel(rows: list, sector: str, as_of: str, out: Path) -> bool:
    try:
        import openpyxl  # type: ignore
        from copy import copy
    except ImportError:
        print("Excel skipped: `pip install openpyxl` to fill templates/p2p_screener_template.xlsx")
        return False
    wb = openpyxl.load_workbook(TEMPLATE_XLSX)
    ws = wb["Screener"]
    ws["A2"] = f"{sector} Sector"
    ws["T2"] = as_of
    title_style = [copy(ws.cell(row=7, column=col)._style) for col in range(1, 23)]
    data_style = [copy(ws.cell(row=11, column=col)._style) for col in range(1, 23)]
    for r in range(7, ws.max_row + 1):
        for col in range(1, 23):
            ws.cell(row=r, column=col).value = None
    columns = {"name": 1, "currency": 3, "price": 4, "vwap90": 5, "mkt_cap_m": 6, "ev_m": 7, "d1": 9, "w1": 10, "ytd": 11,
               "vs_vwap": 12, "ev_sales": 14, "ev_ebitda": 15, "net_leverage": 16, "max_entry_price": 18, "current_takeout": 19,
               "current_irr": 20, "multiple_price": 21}
    r = 7
    for section in SECTIONS:
        block = [row for row in rows if row["section"] == section]
        for col in range(1, 23):
            ws.cell(row=r, column=col)._style = copy(title_style[col - 1])
        ws.cell(row=r, column=1).value = section
        r += 1
        if not block:
            ws.cell(row=r, column=1).value = "(none)"
            r += 1
        for row in block:
            for col in range(1, 23):
                ws.cell(row=r, column=col)._style = copy(data_style[col - 1])
            for key, col in columns.items():
                ws.cell(row=r, column=col).value = row.get(key)
            r += 1
        r += 1
    if "Sources" in wb.sheetnames:
        src = wb["Sources"]
        for rr in range(2, src.max_row + 1):
            for col in range(1, 8):
                src.cell(row=rr, column=col).value = None
        rr = 2
        for row in rows:
            for key, unit in (("price", "local / share"), ("mkt_cap_m", "local, m"), ("ev_m", "local, m"),
                              ("net_leverage", "x"), ("max_entry_price", "local / share")):
                src.append([row["name"], row["ticker"], key, row.get(key), unit, "no", row["src"].get("target_price" if key == "max_entry_price" else "price" if key == "price" else "financials")])
                rr += 1
    wb.save(out)
    return True


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--thesis", type=Path, default=ROOT / "thesis.md")
    ap.add_argument("--out", type=Path, default=ROOT / "outputs")
    ap.add_argument("--excel", action="store_true", help="also fill the Excel screener (needs openpyxl)")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    rows, rejected, failed, settings = build_rows(args.thesis, args.out)
    as_of = json.loads((ROOT / "data" / "market.json").read_text(encoding="utf-8"))["as_of"]
    sector = sector_name(args.thesis)
    write_html(rows, sector, as_of, args.out / "p2p_dashboard.html")
    write_csv(rows, args.out / "screener.csv")
    counts = {s: sum(r["section"] == s for r in rows) for s in SECTIONS}
    print(f"{sector}, as of {as_of}, target price method {settings.method}: "
          + ", ".join(f"{n} {s}" for s, n in counts.items() if n) + (f", {len(rejected)} rejected (not shown)" if rejected else ""))
    for r in rows:
        if r["section"].startswith("Priority"):
            print(f"  {r['section']:<26} {r['name']:<34} price {r['price']} vs actionable below {r['actionable_below']}")
    if failed:
        print(f"  ! the judge failed on {', '.join(failed)} (shown as Radar): rerun python p2p/judge.py --radar")
    print(f"Wrote {args.out / 'p2p_dashboard.html'} and screener.csv")
    if args.excel and write_excel(rows, sector, as_of, args.out / "p2p_screener.xlsx"):
        print(f"Wrote {args.out / 'p2p_screener.xlsx'}")


if __name__ == "__main__":
    main()
